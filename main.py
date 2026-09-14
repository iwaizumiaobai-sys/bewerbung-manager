"""Bewerbungs-Manager — alles in einer Datei.

Start lokal:  uvicorn main:app --reload
Auf Railway:  laeuft ueber das Procfile automatisch.
"""

import email
import io
import json
import os
import re
import sqlite3
import time
from datetime import datetime, timedelta, timezone
from email.header import decode_header

from fastapi import FastAPI, File, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- Einstellungen

FIRMA = "Places to Be"
FARBE = "#1B3A8C"          # Hauptfarbe aus dem Logo. Kopfbereich und Knoepfe.
FARBE_HELL = "#E9EFFB"     # Sehr helle Variante fuer Flaechen.
AKZENT = "#D42B27"         # Das Rot aus dem Logo, sparsam eingesetzt.


def abdunkeln(hexwert, anteil=0.35):
    """Macht eine Farbe dunkler, fuer den Verlauf im Kopfbereich."""
    hexwert = hexwert.lstrip("#")
    r, g, b = (int(hexwert[i:i + 2], 16) for i in (0, 2, 4))
    return "#%02x%02x%02x" % tuple(int(k * (1 - anteil)) for k in (r, g, b))


def kuerzel_aus(name):
    """Bildet ein Kuerzel, falls kein Logo hinterlegt ist."""
    teile = [w for w in re.split(r"\s+", name) if w]
    return "".join(w[0] for w in teile[:3]).upper() or "?"
AUFBEWAHRUNG_TAGE = 3

DB = os.environ.get("DB_PATH", "/tmp/bewerbungen.db")
API_KEY = os.environ.get("ANTHROPIC_API_KEY", "")
MODELL = os.environ.get("AI_MODEL", "claude-sonnet-4-6")

KRITERIEN = [
    ("resilienz", "Umgang mit Ablehnung", 25),
    ("reise", "Reisebereitschaft", 20),
    ("extro", "Extrovertiertheit", 20),
    ("erfahrung", "Vorerfahrung", 15),
    ("motivation", "Motivation", 15),
    ("sprache", "Sprachkenntnisse", 5),
]

STATUS_FARBEN = {
    "gruen": ("C6EFCE", "Einladen"),
    "gelb": ("FFEB9C", "Pruefen"),
    "grau": ("E7E6E6", "Angaben fehlen"),
    "rot": ("FFC7CE", "Unpassend"),
}

# ---------------------------------------------------------------- Datenbank


def conn():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    return c


def init():
    with conn() as c:
        c.execute("""
            CREATE TABLE IF NOT EXISTS bewerbungen (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT, absender TEXT, betreff TEXT, text TEXT,
                status TEXT, gesamt REAL, scores TEXT,
                zusammenfassung TEXT, fuehrerschein TEXT,
                angelegt TEXT
            )
        """)


def aufraeumen():
    grenze = (datetime.now(timezone.utc) - timedelta(days=AUFBEWAHRUNG_TAGE)).isoformat()
    with conn() as c:
        c.execute("DELETE FROM bewerbungen WHERE angelegt < ?", (grenze,))


# ---------------------------------------------------------------- Mail lesen


def kopf(wert):
    if not wert:
        return ""
    teile = []
    for stueck, kodierung in decode_header(wert):
        if isinstance(stueck, bytes):
            teile.append(stueck.decode(kodierung or "utf-8", errors="replace"))
        else:
            teile.append(stueck)
    return "".join(teile).strip()


def html_zu_text(roh):
    roh = re.sub(r"(?is)<(script|style).*?</\1>", " ", roh)
    roh = re.sub(r"(?i)<br\s*/?>", "\n", roh)
    roh = re.sub(r"(?i)</p>", "\n\n", roh)
    roh = re.sub(r"<[^>]+>", " ", roh)
    roh = roh.replace("&nbsp;", " ").replace("&amp;", "&")
    roh = roh.replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"')
    return re.sub(r"\n{3,}", "\n\n", re.sub(r"[ \t]{2,}", " ", roh)).strip()


def pdf_zu_text(rohdaten):
    """Liest den Text aus einer PDF. Gibt leeren String zurueck, wenn es nicht geht."""
    try:
        from pypdf import PdfReader

        leser = PdfReader(io.BytesIO(rohdaten))
        seiten = [(s.extract_text() or "") for s in leser.pages[:15]]
        return re.sub(r"\n{3,}", "\n\n", "\n".join(seiten)).strip()
    except Exception:
        return ""


def eml_lesen(rohdaten):
    msg = email.message_from_bytes(rohdaten)
    absender = kopf(msg.get("From"))
    betreff = kopf(msg.get("Subject"))

    text, html, anhaenge = "", "", []
    if msg.is_multipart():
        for teil in msg.walk():
            if teil.get_content_maintype() == "multipart":
                continue
            dateiname = kopf(teil.get_filename() or "")
            ist_anhang = ("attachment" in str(teil.get("Content-Disposition", ""))
                          or dateiname)
            if ist_anhang:
                if dateiname.lower().endswith(".pdf"):
                    inhalt = teil.get_payload(decode=True)
                    if inhalt:
                        gelesen = pdf_zu_text(inhalt)
                        if gelesen:
                            anhaenge.append(f"[Anhang {dateiname}]\n{gelesen}")
                continue
            try:
                inhalt = teil.get_payload(decode=True)
                if inhalt is None:
                    continue
                zeichensatz = teil.get_content_charset() or "utf-8"
                entpackt = inhalt.decode(zeichensatz, errors="replace")
            except Exception:
                continue
            if teil.get_content_type() == "text/plain" and not text:
                text = entpackt
            elif teil.get_content_type() == "text/html" and not html:
                html = entpackt
    else:
        inhalt = msg.get_payload(decode=True) or b""
        zeichensatz = msg.get_content_charset() or "utf-8"
        entpackt = inhalt.decode(zeichensatz, errors="replace")
        if msg.get_content_type() == "text/html":
            html = entpackt
        else:
            text = entpackt

    koerper = text.strip() or html_zu_text(html)
    if anhaenge:
        koerper = (koerper + "\n\n" + "\n\n".join(anhaenge)).strip()

    name = absender
    if "<" in absender:
        name = absender.split("<")[0].strip().strip('"')
    if not name:
        name = absender

    return {"name": name or "Unbekannt", "absender": absender,
            "betreff": betreff or "(ohne Betreff)", "text": koerper}


# ---------------------------------------------------------------- Bewertung

def mail_raten(text):
    treffer = re.search(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", text)
    return treffer.group(0).rstrip(".,;:") if treffer else ""


def name_raten(text, dateiname=""):
    """Sucht den Namen: erst in der Grussformel am Ende, dann im Dateinamen."""
    schluss = text[-400:]
    for muster in (r"(?:Viele|Beste|Freundliche|Liebe|Herzliche)\s+Gr[uü][sß]{1,2}e,?\s*\n+\s*"
                   r"([A-ZÄÖÜ][\wäöüß'-]+(?:\s+[A-ZÄÖÜ][\wäöüß'-]+){0,2})",
                   r"Mit freundlichen Gr[uü][sß]{1,2}en,?\s*\n+\s*"
                   r"([A-ZÄÖÜ][\wäöüß'-]+(?:\s+[A-ZÄÖÜ][\wäöüß'-]+){0,2})"):
        treffer = re.search(muster, schluss)
        if treffer:
            return treffer.group(1).strip()

    treffer = re.search(r"(?:mein Name ist|Ich hei[sß]e)\s+"
                        r"([A-ZÄÖÜ][\wäöüß'-]+(?:\s+[A-ZÄÖÜ][\wäöüß'-]+){0,2})", text[:1500])
    if treffer:
        return treffer.group(1).strip()

    roh = re.sub(r"\.[^.]+$", "", dateiname)
    roh = re.sub(r"(?i)bewerbung|lebenslauf|anschreiben|cv|final|neu", " ", roh)
    roh = re.sub(r"[_\-.]+", " ", roh).strip()
    woerter = [w for w in roh.split() if len(w) > 1 and not w.isdigit()][:3]
    return " ".join(w.capitalize() for w in woerter) or "Unbekannt"


PROMPT = """Du hilfst einer Fundraising-Agentur bei der Vorsortierung von Bewerbungen
fuer die Taetigkeit als Werber an Infostaenden (Mitgliederwerbung fuer gemeinnuetzige
Organisationen).

Bewerte den folgenden Bewerbungstext auf einer Skala von 1 bis 10 je Kriterium:

- resilienz: Umgang mit Ablehnung, Frustrationstoleranz
- reise: Reisebereitschaft, zeitliche und raeumliche Flexibilitaet
- extro: Offenheit, Kommunikationsfreude, Zugehen auf Fremde
- erfahrung: Vertrieb, Promotion, Standarbeit, Kundenkontakt
- motivation: Warum gerade Fundraising, Bezug zu gemeinnuetziger Arbeit
- sprache: Ausdruck und erwaehnte Sprachkenntnisse

Wenn zu einem Kriterium nichts im Text steht, gib 3 und erwaehne die Luecke.

Antworte ausschliesslich mit JSON, ohne Vorrede und ohne Codebloecke:
{{"resilienz": 0, "reise": 0, "extro": 0, "erfahrung": 0, "motivation": 0,
"sprache": 0, "fuehrerschein": "ja|nein|unbekannt",
"zusammenfassung": "ein bis zwei Saetze"}}

Betreff: {betreff}

Text:
{text}"""


def bewerten(betreff, text):
    if not API_KEY:
        return {k: 5 for k, _, _ in KRITERIEN} | {
            "fuehrerschein": "unbekannt",
            "zusammenfassung": "Keine KI-Bewertung aktiv (ANTHROPIC_API_KEY fehlt).",
        }
    try:
        import anthropic

        klient = anthropic.Anthropic(api_key=API_KEY)
        antwort = klient.messages.create(
            model=MODELL,
            max_tokens=700,
            messages=[{"role": "user",
                       "content": PROMPT.format(betreff=betreff, text=text[:12000])}],
        )
        roh = "".join(b.text for b in antwort.content if b.type == "text")
        roh = re.sub(r"```(?:json)?|```", "", roh).strip()
        daten = json.loads(roh)
    except Exception as fehler:
        return {k: 3 for k, _, _ in KRITERIEN} | {
            "fuehrerschein": "unbekannt",
            "zusammenfassung": f"Bewertung fehlgeschlagen: {fehler}",
        }

    ergebnis = {}
    for schluessel, _, _ in KRITERIEN:
        try:
            ergebnis[schluessel] = max(1, min(10, int(round(float(daten.get(schluessel, 3))))))
        except Exception:
            ergebnis[schluessel] = 3
    ergebnis["fuehrerschein"] = str(daten.get("fuehrerschein", "unbekannt"))[:20]
    ergebnis["zusammenfassung"] = str(daten.get("zusammenfassung", ""))[:500]
    return ergebnis


def gesamtnote(scores):
    summe = sum(scores[k] * g for k, _, g in KRITERIEN)
    return round(summe / 100, 1)


def status_aus(note, scores):
    fehlend = sum(1 for k, _, _ in KRITERIEN if scores[k] == 3)
    if fehlend >= 3:
        return "grau"
    if note >= 7.5:
        return "gruen"
    if note >= 5:
        return "gelb"
    return "rot"


def verarbeiten(daten):
    scores = bewerten(daten["betreff"], daten["text"])
    note = gesamtnote(scores)
    status = status_aus(note, scores)
    nur_scores = {k: scores[k] for k, _, _ in KRITERIEN}

    with conn() as c:
        zeiger = c.execute(
            """INSERT INTO bewerbungen
               (name, absender, betreff, text, status, gesamt, scores,
                zusammenfassung, fuehrerschein, angelegt)
               VALUES (?,?,?,?,?,?,?,?,?,?)""",
            (daten["name"], daten["absender"], daten["betreff"], daten["text"],
             status, note, json.dumps(nur_scores), scores["zusammenfassung"],
             scores["fuehrerschein"], datetime.now(timezone.utc).isoformat()),
        )
        neue_id = zeiger.lastrowid

    return {"id": neue_id, "name": daten["name"], "betreff": daten["betreff"],
            "status": status, "gesamt": note, "scores": nur_scores,
            "fuehrerschein": scores["fuehrerschein"],
            "zusammenfassung": scores["zusammenfassung"]}


# ---------------------------------------------------------------- Web

app = FastAPI()
init()


@app.get("/", response_class=HTMLResponse)
def seite():
    aufraeumen()
    balken = "".join(
        f'<div class="zeile"><div class="lab"><span>{titel}'
        f'<em>{gewicht}%</em></span><b id="w-{key}">–</b></div>'
        f'<div class="spur"><i id="b-{key}"></i></div></div>'
        for key, titel, gewicht in KRITERIEN
    )
    return (SEITE.replace("{{FIRMA}}", FIRMA)
                 .replace("{{KUERZEL}}", kuerzel_aus(FIRMA))
                 .replace("{{FARBE_DUNKEL}}", abdunkeln(FARBE))
                 .replace("{{FARBE_HELL}}", FARBE_HELL)
                 .replace("{{AKZENT}}", AKZENT)
                 .replace("{{FARBE}}", FARBE)
                 .replace("{{BALKEN}}", balken)
                 .replace("{{TAGE}}", str(AUFBEWAHRUNG_TAGE)))


LOGO_URL = os.environ.get("LOGO_URL", "")   # Falls kein logo.png im Repo liegt.


@app.get("/static/logo.png")
def logo():
    from fastapi.responses import FileResponse, RedirectResponse

    if os.path.exists("logo.png"):
        return FileResponse("logo.png")
    if LOGO_URL:
        return RedirectResponse(LOGO_URL)
    return JSONResponse({"fehler": "kein Logo hinterlegt"}, 404)


@app.post("/api/upload")
async def upload(file: UploadFile = File(...)):
    rohdaten = await file.read()
    dateiname = file.filename or ""
    endung = dateiname.lower().rsplit(".", 1)[-1] if "." in dateiname else ""

    if endung == "eml":
        try:
            daten = eml_lesen(rohdaten)
        except Exception as fehler:
            return JSONResponse({"fehler": f"Datei nicht lesbar: {fehler}"}, 400)

    elif endung == "pdf":
        text = pdf_zu_text(rohdaten)
        if not text:
            return JSONResponse(
                {"fehler": "Aus dieser PDF laesst sich kein Text lesen. "
                           "Vermutlich ein Scan — bitte den Text von Hand einfuegen."}, 400)
        daten = {"name": name_raten(text, dateiname), "absender": mail_raten(text),
                 "betreff": f"Bewerbung ({dateiname})", "text": text}

    elif endung in ("txt", "text", "md", "rtf"):
        text = rohdaten.decode("utf-8", errors="replace")
        daten = {"name": name_raten(text, dateiname), "absender": mail_raten(text),
                 "betreff": f"Bewerbung ({dateiname})", "text": text}

    else:
        return JSONResponse(
            {"fehler": "Moegliche Formate: .eml, .pdf, .txt"}, 400)

    if not daten["text"].strip():
        return JSONResponse({"fehler": "Die Datei enthaelt keinen Text."}, 400)
    return verarbeiten(daten)


@app.post("/api/text")
async def per_hand(nutzlast: dict):
    text = (nutzlast.get("text") or "").strip()
    if not text:
        return JSONResponse({"fehler": "Kein Text angegeben."}, 400)
    return verarbeiten({
        "name": (nutzlast.get("name") or "Unbekannt").strip(),
        "absender": (nutzlast.get("absender") or "").strip(),
        "betreff": (nutzlast.get("betreff") or "Bewerbung").strip(),
        "text": text,
    })


@app.get("/mail/{eintrag_id}", response_class=HTMLResponse)
def mail_ansehen(eintrag_id: int):
    aufraeumen()
    with conn() as c:
        reihe = c.execute("SELECT * FROM bewerbungen WHERE id=?", (eintrag_id,)).fetchone()
    if reihe is None:
        return HTMLResponse(
            f"<body style='font-family:system-ui;padding:40px;max-width:600px;margin:auto'>"
            f"<h2>Nicht mehr vorhanden</h2><p>Bewerbungen werden nach "
            f"{AUFBEWAHRUNG_TAGE} Tagen automatisch geloescht.</p></body>", 404)
    sicher = (reihe["text"].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
    return HTMLResponse(f"""<!doctype html><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<body style="font-family:system-ui,-apple-system,sans-serif;max-width:760px;margin:0 auto;padding:28px;color:#1c1c1a">
<div style="background:{FARBE};color:#fff;padding:16px 20px;border-radius:12px">
<div style="font-size:13px;opacity:.85">{reihe['betreff']}</div>
<div style="font-size:20px;font-weight:600;margin-top:2px">{reihe['name']}</div></div>
<p style="color:#6b6963;font-size:13px;margin:14px 0 4px">{reihe['absender']}</p>
<p style="color:#6b6963;font-size:13px;margin:0 0 18px">Gesamtnote {reihe['gesamt']}/10</p>
<pre style="white-space:pre-wrap;font:inherit;line-height:1.65;background:#faf9f7;border:1px solid #e6e3dd;border-radius:12px;padding:18px">{sicher}</pre>
</body>""")


@app.get("/api/excel")
def excel():
    aufraeumen()
    with conn() as c:
        reihen = c.execute(
            "SELECT * FROM bewerbungen ORDER BY gesamt DESC, id DESC").fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Bewerbungen"

    kopfzeile = (["Status", "Name", "Absender", "Gesamt (1-10)"]
                 + [t for _, t, _ in KRITERIEN]
                 + ["Fuehrerschein", "Zusammenfassung", "Einladung", "Mail"])

    rand = Border(*[Side(style="thin", color="BFBFBF")] * 4)
    for spalte, titel in enumerate(kopfzeile, 1):
        zelle = ws.cell(row=1, column=spalte, value=titel)
        zelle.fill = PatternFill("solid", start_color="2F4F4F")
        zelle.font = Font(bold=True, color="FFFFFF", size=10)
        zelle.alignment = Alignment("center", "center", wrap_text=True)
        zelle.border = rand
    ws.row_dimensions[1].height = 30

    basis = os.environ.get("BASIS_URL", "").rstrip("/")

    for nr, reihe in enumerate(reihen, start=2):
        scores = json.loads(reihe["scores"])
        farbe, beschriftung = STATUS_FARBEN.get(reihe["status"], ("FFFFFF", ""))
        werte = ([beschriftung, reihe["name"], reihe["absender"], reihe["gesamt"]]
                 + [scores.get(k, "") for k, _, _ in KRITERIEN]
                 + [reihe["fuehrerschein"], reihe["zusammenfassung"], ""])

        for spalte, wert in enumerate(werte, 1):
            zelle = ws.cell(row=nr, column=spalte, value=wert)
            zelle.fill = PatternFill("solid", start_color=farbe)
            zelle.border = rand
            zelle.alignment = Alignment(wrap_text=True, vertical="top")

        link = ws.cell(row=nr, column=len(kopfzeile), value="Oeffnen")
        link.hyperlink = f"{basis}/mail/{reihe['id']}"
        link.font = Font(color="0563C1", underline="single")
        link.fill = PatternFill("solid", start_color=farbe)
        link.border = rand
        link.alignment = Alignment("center", "center")
        ws.row_dimensions[nr].height = 44

    breiten = [14, 22, 26, 12] + [11] * len(KRITERIEN) + [13, 46, 11, 10]
    for i, breite in enumerate(breiten, 1):
        ws.column_dimensions[get_column_letter(i)].width = breite
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(kopfzeile))}{max(1, len(reihen) + 1)}"

    hinweis = len(reihen) + 3
    ws.cell(row=hinweis, column=1,
            value=f"Mail-Links sind {AUFBEWAHRUNG_TAGE} Tage gueltig. "
                  f"Stand: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
            ).font = Font(size=9, italic=True, color="888888")

    puffer = io.BytesIO()
    wb.save(puffer)
    puffer.seek(0)
    return StreamingResponse(
        puffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition":
                 f'attachment; filename="bewerbungen_{datetime.now():%Y-%m-%d}.xlsx"'},
    )


# ---------------------------------------------------------------- Oberflaeche

SEITE = """<!doctype html><html lang="de"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Bewerbungen — {{FIRMA}}</title><style>
*{box-sizing:border-box}
body{margin:0;font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",system-ui,sans-serif;
color:#152238;line-height:1.6;min-height:100vh;
background:
radial-gradient(900px 520px at 88% 96%,rgba(120,170,235,.20),transparent 62%),
radial-gradient(700px 420px at 4% 4%,rgba(150,195,245,.16),transparent 60%),
linear-gradient(170deg,#f4f8fe 0%,#e9f1fc 100%)}
.huelle{max-width:720px;margin:0 auto;padding:18px 16px 40px}

/* Kopfzeile mit Logo */
.marke{display:flex;align-items:center;gap:12px;padding:4px 2px 0}
.marke img{height:46px;width:auto}
.marke .kuerzel{height:46px;width:46px;flex:0 0 46px;border-radius:50%;
background:linear-gradient(140deg,{{FARBE}},{{FARBE_DUNKEL}});color:#fff;
display:flex;align-items:center;justify-content:center;font-size:15px;font-weight:700}
.marke h1{margin:0;font-size:21px;font-weight:700;letter-spacing:-.02em;color:#12224a}
.marke h1 i{font-style:normal;color:{{AKZENT}}}
.marke p{margin:0;font-size:11px;letter-spacing:.13em;text-transform:uppercase;
color:#8fa2c2;font-weight:500}

/* Blaues Band */
.band{margin-top:18px;border-radius:18px;padding:22px 24px;color:#fff;position:relative;
overflow:hidden;display:flex;align-items:center;gap:16px;
background:linear-gradient(115deg,{{FARBE_DUNKEL}} 0%,{{FARBE}} 58%,#2a5fd0 100%);
box-shadow:0 10px 30px rgba(20,50,120,.26)}
.band:before{content:"";position:absolute;right:-90px;top:-70px;width:280px;height:280px;
border-radius:50%;background:radial-gradient(circle at 34% 34%,
rgba(120,190,255,.55),rgba(60,120,220,.20) 52%,transparent 72%)}
.band .ikon{width:52px;height:52px;flex:0 0 52px;border-radius:14px;
background:rgba(255,255,255,.16);display:flex;align-items:center;justify-content:center;
font-size:24px;position:relative;z-index:1}
.band .txt{position:relative;z-index:1;min-width:0}
.band h2{margin:0;font-size:19px;font-weight:700;letter-spacing:-.015em}
.band p{margin:2px 0 0;font-size:13.5px;opacity:.86}
.band .strich{position:relative;z-index:1;margin-top:12px;height:4px;width:min(300px,60%);
border-radius:3px;background:rgba(255,255,255,.22);overflow:hidden}
.band .strich i{display:block;height:100%;width:62%;border-radius:3px;
background:linear-gradient(90deg,#7fc0ff,#d8ecff)}

/* Ablageflaeche */
.zone{margin-top:16px;background:#fff;border-radius:20px;padding:12px;
box-shadow:0 6px 26px rgba(30,70,150,.10);transition:.18s ease}
.zone .innen{border:2px dashed #b9cdea;border-radius:15px;padding:36px 18px;
text-align:center;transition:.18s ease}
.zone.aktiv{transform:scale(1.012)}
.zone.aktiv .innen{border-color:{{FARBE}};background:{{FARBE_HELL}}}
.zone .sym{width:66px;height:66px;margin:0 auto;border-radius:50%;
background:radial-gradient(circle at 50% 38%,#fff,{{FARBE_HELL}});
border:1px solid #dbe7f8;display:flex;align-items:center;justify-content:center;
font-size:28px;box-shadow:0 3px 14px rgba(40,90,180,.13)}
.zone h3{margin:14px 0 2px;font-size:21px;font-weight:700;letter-spacing:-.02em;color:#12224a}
.zone .unter{margin:0;color:#7e8ca6;font-size:14px}
.marken{display:flex;gap:8px;justify-content:center;margin-top:13px;flex-wrap:wrap}
.marken span{font-size:12px;font-weight:700;letter-spacing:.5px;color:{{FARBE}};
background:{{FARBE_HELL}};border:1px solid #dbe7f8;border-radius:9px;padding:5px 15px}

/* Knoepfe */
button{font-family:inherit;cursor:pointer;border:0;font-weight:600}
.haupt{margin-top:18px;background:linear-gradient(100deg,{{FARBE}},#2f6ae0);
color:#fff;border-radius:13px;padding:15px 26px;font-size:15.5px;width:100%;max-width:330px;
display:inline-flex;align-items:center;justify-content:center;gap:10px;
box-shadow:0 7px 20px rgba(30,80,190,.30);transition:.15s}
.haupt:active{transform:translateY(1px);box-shadow:0 4px 12px rgba(30,80,190,.28)}
.haupt .pfeil{font-size:18px;line-height:1}
.zweit{width:100%;background:#fff;color:{{FARBE}};border:1.6px solid {{FARBE}};
border-radius:13px;padding:14px 20px;font-size:15px;
display:inline-flex;align-items:center;justify-content:center;gap:10px}
.zweit:active{background:{{FARBE_HELL}}}

.trenner{display:flex;align-items:center;gap:12px;margin:22px 0 14px;
color:#8fa2c2;font-size:13px;font-weight:600}
.trenner:before,.trenner:after{content:"";flex:1;height:1px;background:#d6e2f3}

/* Karten */
.karte{background:#fff;border-radius:20px;padding:20px 22px;margin-top:16px;
box-shadow:0 6px 26px rgba(30,70,150,.10)}
.karte.aus{display:none}
input,textarea{width:100%;padding:12px 14px;border:1px solid #d5e0f2;border-radius:11px;
font-size:15px;font-family:inherit;margin-bottom:10px;background:#fbfcff;color:#152238}
input:focus,textarea:focus{outline:0;border-color:{{FARBE}};background:#fff}
textarea{min-height:150px;resize:vertical}

.oben{display:flex;justify-content:space-between;align-items:flex-start;gap:12px}
.nam{font-size:18px;font-weight:700;margin:0;letter-spacing:-.01em;color:#12224a}
.pill{display:inline-block;font-size:12px;font-weight:700;padding:4px 12px;
border-radius:20px;margin-top:7px}
.note{font-size:36px;font-weight:700;line-height:1;text-align:right;color:#12224a}
.note span{font-size:15px;color:#9aa9c2;font-weight:500}
.note small{display:block;font-size:12px;color:#7e8ca6;font-weight:500;margin-bottom:2px}
hr{border:0;border-top:1px solid #e7eefa;margin:16px 0}
.zeile{margin-bottom:12px}
.lab{display:flex;justify-content:space-between;font-size:13px;margin-bottom:5px}
.lab span{color:#7e8ca6}
.lab em{font-style:normal;color:#a7b5cc;font-size:11px;margin-left:5px}
.lab b{font-weight:700;color:#12224a}
.spur{height:9px;background:#eef3fb;border-radius:5px;overflow:hidden}
.spur i{display:block;height:100%;width:0;border-radius:5px;transition:width .45s ease}
.fazit{font-size:14px;color:#48566e;margin:0}
.fuss{font-size:13px;color:#7e8ca6;margin:6px 0 0}

.lade{text-align:center;color:{{FARBE}};font-size:14px;font-weight:600;padding:16px 0}
.fehler{background:#fdecea;color:{{AKZENT}};border-radius:12px;padding:12px 15px;
font-size:14px;margin-top:12px;font-weight:500}

.sicher{display:flex;align-items:center;justify-content:center;gap:9px;margin:26px 0 0}
.sicher .schild{font-size:19px}
.sicher p{margin:0;font-size:13px;color:#5b6b86}
.sicher p b{color:#12224a}
.sicher p small{display:block;font-size:11.5px;color:#93a3bf}
</style></head><body><div class="huelle">

<div class="marke">
<img src="/static/logo.png" alt="" onerror="this.outerHTML='<div class=\\'kuerzel\\'>{{KUERZEL}}</div>'">
<div><h1>Places <i>to</i> Be</h1><p>Bewerbungen sichten</p></div>
</div>

<div class="band">
<div class="ikon">📄</div>
<div class="txt"><h2>Bewerbungen verarbeiten</h2>
<p>Unterlagen hochladen und bewerten lassen.</p>
<div class="strich"><i></i></div></div>
</div>

<div class="zone" id="zone"><div class="innen">
<div class="sym">☁️</div>
<h3>Datei hier ablegen</h3>
<p class="unter">oder unten auswählen</p>
<div class="marken"><span>EML</span><span>PDF</span><span>TXT</span></div>
<button class="haupt" onclick="datei.click()">Datei auswählen <span class="pfeil">→</span></button>
<input type="file" id="datei" accept=".eml,.pdf,.txt,.md,.rtf" hidden>
</div></div>

<div class="trenner">oder</div>
<button class="zweit" onclick="handForm()">⌨️ Text von Hand eingeben</button>

<div class="karte aus" id="hand">
<input id="h-name" placeholder="Name">
<input id="h-mail" placeholder="E-Mail-Adresse">
<input id="h-betreff" placeholder="Betreff">
<textarea id="h-text" placeholder="Bewerbungstext einfügen"></textarea>
<button class="haupt" style="max-width:none" onclick="sendeText()">Bewerten <span class="pfeil">→</span></button>
</div>

<div class="lade" id="lade" style="display:none">Wird ausgewertet …</div>
<div id="fehler"></div>

<div class="karte aus" id="ergebnis">
<div class="oben">
<div><p class="nam" id="r-name">–</p><span class="pill" id="r-pill">–</span></div>
<div class="note"><small>Gesamt</small><span id="r-note">–</span><span>/10</span></div>
</div>
<hr>
{{BALKEN}}
<hr>
<p class="fazit" id="r-fazit">–</p>
<p class="fuss" id="r-fs">–</p>
<button class="haupt" style="max-width:none" onclick="location.href='/api/excel'">
⬇ Excel herunterladen</button>
<button class="zweit" style="margin-top:9px" onclick="zuruecksetzen()">
Nächste Bewerbung</button>
</div>

<div class="sicher"><span class="schild">🛡️</span>
<p><b>Die Entscheidung trifft das Team.</b>
<small>Bewerbungen werden nach {{TAGE}} Tagen automatisch gelöscht.</small></p></div>
</div>

<script>
const zone=document.getElementById('zone'),datei=document.getElementById('datei');
const KRIT=[{{KRITLISTE}}];

['dragenter','dragover'].forEach(e=>zone.addEventListener(e,v=>{
v.preventDefault();zone.classList.add('aktiv')}));
['dragleave','drop'].forEach(e=>zone.addEventListener(e,v=>{
v.preventDefault();zone.classList.remove('aktiv')}));
zone.addEventListener('drop',v=>{if(v.dataTransfer.files[0])schicke(v.dataTransfer.files[0])});
datei.addEventListener('change',()=>{if(datei.files[0])schicke(datei.files[0])});

function handForm(){document.getElementById('hand').classList.toggle('aus')}
function laden(an){document.getElementById('lade').style.display=an?'block':'none'}
function fehler(t){document.getElementById('fehler').innerHTML=
t?'<div class="fehler">'+t+'</div>':''}

async function schicke(f){
fehler('');laden(true);
const fd=new FormData();fd.append('file',f);
try{const a=await fetch('/api/upload',{method:'POST',body:fd});
const d=await a.json();if(!a.ok)throw new Error(d.fehler||'Fehler');zeige(d)}
catch(e){fehler(e.message)}finally{laden(false);datei.value=''}}

async function sendeText(){
const text=document.getElementById('h-text').value.trim();
if(!text){fehler('Bitte den Bewerbungstext einfügen.');return}
fehler('');laden(true);
try{const a=await fetch('/api/text',{method:'POST',
headers:{'Content-Type':'application/json'},body:JSON.stringify({
name:document.getElementById('h-name').value,
absender:document.getElementById('h-mail').value,
betreff:document.getElementById('h-betreff').value,text:text})});
const d=await a.json();if(!a.ok)throw new Error(d.fehler||'Fehler');
document.getElementById('hand').classList.add('aus');zeige(d)}
catch(e){fehler(e.message)}finally{laden(false)}}

function farbe(n){return n>=7?'#4f8a3d':n>=4?'#c98a1e':'#c4483f'}
const PILL={gruen:['#d9efdc','#2a6b36','Einladen'],gelb:['#faeecb','#8a6410','Prüfen'],
grau:['#ebe9e3','#5f5c56','Angaben fehlen'],rot:['#f8dcd9','#9c3229','Unpassend']};

function zeige(d){
document.getElementById('r-name').textContent=d.name;
document.getElementById('r-note').textContent=d.gesamt;
const p=PILL[d.status]||PILL.grau,pill=document.getElementById('r-pill');
pill.textContent=p[2];pill.style.background=p[0];pill.style.color=p[1];
KRIT.forEach(k=>{const v=d.scores[k]||0;
document.getElementById('w-'+k).textContent=v;
const b=document.getElementById('b-'+k);
b.style.background=farbe(v);setTimeout(()=>b.style.width=(v*10)+'%',30)});
document.getElementById('r-fazit').textContent=d.zusammenfassung;
document.getElementById('r-fs').textContent='Führerschein: '+d.fuehrerschein;
document.getElementById('ergebnis').classList.remove('aus');
document.getElementById('ergebnis').scrollIntoView({behavior:'smooth',block:'start'})}

function zuruecksetzen(){
document.getElementById('ergebnis').classList.add('aus');
KRIT.forEach(k=>document.getElementById('b-'+k).style.width='0');
['h-name','h-mail','h-betreff','h-text'].forEach(i=>document.getElementById(i).value='');
fehler('');window.scrollTo({top:0,behavior:'smooth'})}
</script></body></html>"""

SEITE = SEITE.replace("{{KRITLISTE}}", ",".join(f"'{k}'" for k, _, _ in KRITERIEN))
